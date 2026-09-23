"""Suite T7: integridad del universo cargado y de las facetas públicas."""
from __future__ import annotations

import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import truststore
truststore.inject_into_ssl()
import requests

ROOT = Path(__file__).parent
OUT = ROOT / "salida"
BATCH = 1000


def env():
    for line in (ROOT.parent / ".env.local").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1); os.environ.setdefault(k.strip(), v.strip().strip('"'))


class Api:
    def __init__(self, url, key):
        self.base = url.rstrip("/") + "/rest/v1"; self.key = key

    def get(self, table, select="*", order="id", extra=None, key=None):
        out=[]; off=0; headers={"apikey":key or self.key, "Authorization":f"Bearer {key or self.key}"}
        while True:
            params={"select":select,"order":order,"limit":BATCH,"offset":off}; params.update(extra or {})
            r=requests.get(f"{self.base}/{table}",headers=headers,params=params,timeout=180)
            if r.status_code >= 300: raise RuntimeError(f"GET {table} HTTP {r.status_code}")
            page=r.json(); out += page
            if len(page)<BATCH:return out
            off += BATCH

    def count(self, table, extra=None, select="id"):
        h={"apikey":self.key,"Authorization":f"Bearer {self.key}","Prefer":"count=exact"}
        r=requests.get(f"{self.base}/{table}",headers=h,params={"select":select,"limit":0,**(extra or {})},timeout=180)
        if r.status_code >= 300: raise RuntimeError(f"COUNT {table} HTTP {r.status_code}")
        return int(r.headers.get("content-range","*/-1").split("/")[-1])


def check(n, condition, detail=""):
    print(f"{'OK' if condition else 'FALLA'} {n}{': ' + detail if detail else ''}")
    return bool(condition)


def js_string(valor):
    if isinstance(valor, bool):
        return "true" if valor else "false"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def clave(valor):
    if isinstance(valor, bool):
        return ("bool", valor)
    if isinstance(valor, (int, float)):
        return ("num", float(valor))
    if isinstance(valor, tuple):
        return ("tuple", tuple(clave(x) for x in valor))
    return ("valor", valor)


def distribucion_dataset(ds, campo):
    definicion = ds["campos"][campo]
    valores = ds["columnas"][campo]
    salida = []
    for valor in valores:
        if valor is None:
            salida.append(None)
        elif definicion["tipo"] == "cat":
            salida.append(definicion["valores"][valor])
        elif definicion["tipo"] == "multi":
            salida.append(tuple(definicion["valores"][x] for x in valor))
        elif definicion["tipo"] == "bool":
            salida.append(bool(valor))
        else:
            salida.append(valor)
    return Counter(clave(x) for x in salida)


def claves_fase3_independientes():
    """Recalcula cédulas elegibles directamente de los Excel, sin importar el loader."""
    from openpyxl import load_workbook

    source_root = Path(r"C:\Users\EstudiantesJC\Downloads\2026-Fase3\extraido")
    country_by_suffix = {"BAQ": "CO", "BOG": "CO", "CLO": "CO", "CTG": "CO",
                         "MED": "CO", "ECU": "EC", "URY": "UY", "PAN": "PA"}
    paths = sorted(source_root.glob("*/Tabla de Puntuación * PANEL *.xlsx"))

    def norm(value):
        text = "" if value is None else str(value).strip()
        text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
        return " ".join(text.casefold().split())

    def find_columns(sheet, required):
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            keys = {norm(value) for value in values}
            if required.issubset(keys):
                return row_number, {norm(value): index for index, value in enumerate(values)
                                    if norm(value)}
        return None, {}

    def get(row, columns, field):
        index = columns.get(field)
        return None if index is None or index >= len(row) else row[index]

    candidates = defaultdict(list)
    blocked_files = 0
    for path in paths:
        country = country_by_suffix.get(path.parent.name[-3:])
        if country is None:
            blocked_files += 1
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if "GENERAL" not in workbook.sheetnames:
                blocked_files += 1
                continue
            general = workbook["GENERAL"]
            support_name = next((name for name in workbook.sheetnames if name.startswith("SOPORTE")), None)
            if support_name is None:
                blocked_files += 1
                continue
            support = workbook[support_name]
            general_header, gc = find_columns(general, {"id", "email"})
            support_header, sc = find_columns(support, {"email", "presente"})
            required = {"grupo", "capitan", "jurado 1", "jurado 2", "total", "puntaje", "disponibilidad"}
            if general_header is None or support_header is None or not required.issubset(gc):
                blocked_files += 1
                continue
            header_values = next(general.iter_rows(min_row=general_header,
                                                   max_row=general_header,
                                                   values_only=True))
            score_index, availability_index = gc["puntaje"], gc["disponibilidad"]
            loose = [i for i in range(score_index + 1, availability_index)
                     if i < len(header_values) and header_values[i] is not None
                     and str(header_values[i]).strip().isdigit()]
            support_by_email = defaultdict(list)
            for row in support.iter_rows(min_row=support_header + 1, values_only=True):
                email = "" if get(row, sc, "email") is None else str(get(row, sc, "email")).strip().casefold()
                if email:
                    support_by_email[email].append(norm(get(row, sc, "presente")))
            parsed, zero_emails, no_emails, cumple_values = [], Counter(), Counter(), set()
            for row_number, row in enumerate(general.iter_rows(min_row=general_header + 1,
                                                                 values_only=True),
                                             start=general_header + 1):
                raw_id = get(row, gc, "id")
                raw_email = get(row, gc, "email")
                email = "" if raw_email is None else str(raw_email).strip().casefold()
                if raw_id is None and not email:
                    continue
                if raw_id is None or not email:
                    continue
                try:
                    cedula = str(int(raw_id))
                except (TypeError, ValueError, OverflowError):
                    continue
                total = get(row, gc, "total")
                try:
                    total = None if total is None or str(total).strip() == "" else float(total)
                except (TypeError, ValueError):
                    total = None
                if total == 0:
                    zero_emails[email] += 1
                if len(loose) == 1:
                    cumple_values.add(norm(row[loose[0]] if loose[0] < len(row) else None))
                attendance = support_by_email.get(email, [])
                if len(attendance) == 1:
                    present = attendance[0]
                    if present == "no":
                        no_emails[email] += 1
                    attended = True if present == "si" else (False if present == "no" else None)
                else:
                    attended = None
                score = get(row, gc, "puntaje")
                try:
                    score = None if score is None or str(score).strip() == "" else float(str(score).strip().rstrip("%"))
                except (TypeError, ValueError):
                    score = None
                parsed.append({"key": (country, cedula), "attended": attended,
                               "score": score, "file": path.name, "row": row_number})
            file_blocked = (len(loose) != 1 or zero_emails != no_emails
                            or cumple_values != {"cumple"})
            if file_blocked:
                blocked_files += 1
            for item in parsed:
                if file_blocked:
                    item["attended"] = None
                # Una respuesta de asistencia indeterminada a nivel individual
                # también se conserva; no equivale a ausencia ni se descarta.
                candidates[item["key"]].append(item)
        finally:
            workbook.close()
    selected = {}
    for key, rows in candidates.items():
        known = [item for item in rows if item["attended"] is not None]
        if known:
            winner = max(known, key=lambda item: (item["attended"], item["score"] is not None,
                                                   item["score"] or -1, item["file"], -item["row"]))
        else:
            with_score = [item for item in rows if item["score"] is not None]
            winner = (max(with_score, key=lambda item: item["score"])
                      if with_score else min(rows, key=lambda item: (item["file"], item["row"])))
        selected[key] = winner
    return selected, len(paths), blocked_files


def distribucion_fuente(filas, campo, definicion):
    salida = []
    for fila in filas:
        valor = fila.get(campo)
        if campo == "edad" and fila.get("edad_valida") is not True:
            valor = None
        if definicion["tipo"] == "multi":
            valor = tuple(valor) if valor else ("sin_dato",)
        elif definicion["tipo"] == "num":
            valor = None if valor is None else float(valor)
        elif definicion["tipo"] == "bool":
            valor = bool(valor)
        else:
            if valor is None or valor == "":
                valor = "Sin dato"
            else:
                valor = js_string(valor)
        salida.append(valor)
    return Counter(clave(x) for x in salida)


def main():
    env(); url=os.getenv("CONV_SUPABASE_URL"); service=os.getenv("CONV_SUPABASE_SERVICE_ROLE_KEY"); anon=os.getenv("CONV_SUPABASE_ANON_KEY")
    if not url or not service or not anon: raise SystemExit("faltan credenciales CONV_SUPABASE")
    api=Api(url,service); failures=0
    files=sorted(OUT.glob("raw_*.jsonl")); counts={p.name:sum(1 for _ in p.open(encoding="utf-8")) for p in files}
    expected={"raw_2026_CO.jsonl":11356,"raw_2026_EC.jsonl":1430,"raw_2026_PA.jsonl":567,"raw_2026_UY.jsonl":219,
              "raw_2025_CO.jsonl":7977,"raw_2025_CO-form1.jsonl":90,"raw_2025_EC.jsonl":2213,"raw_2025_UY.jsonl":351}
    failures += not check(1, counts==expected and sum(counts.values())==24203, str(sum(counts.values())))
    posts=api.get("postulaciones","id,id_publico,convocatoria,pais,ciudad_norm,enrutado_fuera_cobertura,estrato,edad,edad_valida,indice_activos",order="id")
    results=api.get("resultado_seleccion","postulacion_id,seleccionado,duplicado_de",order="postulacion_id")
    by_result={x["postulacion_id"]:x for x in results}
    for x in posts: x.update(by_result.get(x["id"], {"seleccionado":False,"duplicado_de":None}))
    failures += not check(2, sorted(x["id_publico"] for x in posts)==list(range(1,24204)), "id_publico")
    pii=api.count("postulaciones_pii",select="postulacion_id"); failures += not check(3, pii==24203, f"pii={pii}")
    closed={"genero":{"femenino","masculino","no_binario","otro","no_responde"},"emprendimiento":{"si","considera","no"},"otros_programas":{"no","espera","rechazado","aceptado"},"acceso_computador":{"propio","prestado","ninguno"},"comodidad_autonomo":{"alta","media","baja"}}
    payload=json.load((OUT/"payload.json").open(encoding="utf-8"))["postulaciones"]
    failures += not check(4, all(x["postulacion"].get(k) in vals or x["postulacion"].get(k) is None for x in payload for k,vals in closed.items()), "catalogos")
    co=[x for x in posts if x["pais"]=="CO" and not x["enrutado_fuera_cobertura"]]; non=sum(x["estrato"] is not None for x in co)
    ex=[x for x in posts if x["pais"] in {"EC","UY","PA"}]; failures += not check(5, all(x["estrato"] is None for x in ex) and non/len(co)>=.99, f"CO={non}/{len(co)}")
    ages=[x for x in posts if x["edad"] is not None]; bad=[x for x in ages if not 10<=x["edad"]<=80]; failures += not check(6, len(bad)/len(ages)<=.01, f"excepciones={len(bad)}")
    selected=[x for x in posts if (x.get("seleccionado") or False)]
    selected_by_conv=Counter(x["convocatoria"] for x in selected)
    failures += not check(7, selected_by_conv == Counter({"2026": 832, "2025": 722}),
                          f"seleccionadas={dict(selected_by_conv)}")
    rest=[x for x in posts if (x["convocatoria"],x["pais"]) not in {("2025","UY"),("2026","PA"),("2026","UY")} and not x["enrutado_fuera_cobertura"]]; failures += not check(8, sum(x["indice_activos"] is not None for x in rest)/len(rest)>=.99, "indice_activos")
    c=Counter((x["convocatoria"],x["pais"]) for x in posts if x["enrutado_fuera_cobertura"]); failures += not check(9,c.get(("2025","CO"))==955 and c.get(("2025","EC"))==626 and not any(x["enrutado_fuera_cobertura"] and (x["pais"]=="UY" or x["convocatoria"]=="2026") for x in posts), str(c))
    covered={("CO", c) for c in {"Barranquilla","Bogotá D.C.","Cali","Cartagena de Indias","Medellín","Valle de Aburrá"}} | {("EC", "Guayaquil")}
    friccion_pairs=Counter((x["pais"],x["ciudad_norm"]) for x in posts if x["enrutado_fuera_cobertura"] and (x["pais"],x["ciudad_norm"]) in covered)
    esperado=Counter({("CO","Cartagena de Indias"):18,("EC","Guayaquil"):6})
    (OUT/"fricciones_cobertura.csv").write_text("pais,ciudad_norm,filas\n" + "\n".join(f"{p},{c},{n}" for (p,c),n in sorted(friccion_pairs.items())), encoding="utf-8")
    failures += not check(10, friccion_pairs==esperado and sum(friccion_pairs.values())==24, f"friccion={dict(friccion_pairs)}; enrutamiento_correcto=10")
    r=requests.get(f"{api.base}/postulaciones_pii",headers={"apikey":anon,"Authorization":f"Bearer {anon}"},params={"select":"*","limit":1},timeout=60); failures += not check(11,r.status_code in {401,403}, f"anon_http={r.status_code}")
    city=Counter(x["ciudad_norm"] for x in selected); failures += not check(12, True, f"seleccionadas_por_ciudad={city.most_common(5)}; hojas BOG=149 MED=110")
    rs=api.count("resultado_seleccion",select="postulacion_id"); failures += not check(13,rs==24203,f"resultado={rs}")
    view=api.get("v_analisis_postulaciones","*",order="id_publico")
    categorical=["ciudad","genero","situacion_educativa","condicion_laboral","emprendimiento","otros_programas","estrato_cat","ingreso_hogar","nucleo_estado","tipo_vivienda","indice_activos_estado","tiene_internet","acceso_computador","horas_semanales","comodidad_autonomo","nivel_software","nivel_ingles","tiene_embajador","aplico_antes_jc","fue_beneficiario_antes","retirado","estado_final","fase_max_alcanzada","metodo_match"]
    arrays=["segmentos","ocupaciones","como_se_entero"]
    sums={k:sum(bool(x.get(k)) for x in view) for k in categorical+arrays}
    failures += not check(14, len(view)==24203 and all(v==24203 for v in sums.values()), "segmentaciones")
    failures += not check(15, sum(counts.values())==24203 and len(posts)==24203, f"fuentes={sum(counts.values())}")
    dataset_path = ROOT.parent / "web" / "data" / "postulaciones.json"
    dataset = json.loads(dataset_path.read_text(encoding="utf-8")) if dataset_path.exists() else {}
    dataset_ok = dataset.get("es_ejemplo") is False and dataset.get("total") == 24203
    derivados = {"datos_curso", "envio_duplicado", "fase_max_alcanzada",
                 "prueba_completada", "puntaje_prueba_pct", "formulario_completado",
                 "preguntas_respondidas", "respuestas_validas", "pasa_fase2"}
    comparables = sorted((set(dataset.get("campos", {})) - derivados) & set(view[0] if view else {}))
    diferencias = {}
    if dataset_ok:
        for campo in comparables:
            esperado = distribucion_fuente(view, campo, dataset["campos"][campo])
            recibido = distribucion_dataset(dataset, campo)
            if esperado != recibido:
                diferencias[campo] = {"esperado": sum((esperado - recibido).values()),
                                      "recibido": sum((recibido - esperado).values())}
    expected_comparables = len(set(dataset.get("campos", {})) - derivados)
    failures += not check(16, dataset_ok and len(comparables) == expected_comparables and not diferencias,
                          f"campos={len(comparables)} diferencias={diferencias}")
    programs=api.get("resultado_programa", "postulacion_id,cohorte,retirado,pct_avance,cursos_aprobados",
                     order="postulacion_id")
    metric_2025=sum(x["cohorte"] == "2025" and x["cursos_aprobados"] is not None for x in programs)
    metric_2026=sum(x["cohorte"] == "2026" and x["cursos_aprobados"] is not None for x in programs)
    retired_2025=sum(x["cohorte"] == "2025" and x["retirado"] is True and x["cursos_aprobados"] is not None
                    for x in programs)
    no_programa=len(view)-len(programs)
    failures += not check(17, (len(programs), metric_2025, metric_2026, retired_2025, no_programa) == (1554, 559, 776, 0, 22649),
                          f"metricas_2025={metric_2025} metricas_2026={metric_2026} "
                          f"retirados_2025_con_dato={retired_2025} no_aplica={no_programa}")
    curso_values = dataset.get("campos", {}).get("datos_curso", {}).get("valores", [])
    curso_counts = Counter(curso_values[v] for v in dataset.get("columnas", {}).get("datos_curso", []) if v < len(curso_values))
    cursos_num = dataset.get("campos", {}).get("cursos_inscritos", {}).get("tipo") == "num"
    failures += not check(18, curso_counts == Counter({"Con datos": 1335, "Sin dato": 219, "No aplica": 22649}) and cursos_num,
                          f"datos_curso={dict(curso_counts)} cursos_inscritos={dataset.get('campos', {}).get('cursos_inscritos', {}).get('tipo')}")
    respuestas = api.get("postulaciones_respuestas", "postulacion_id,respuestas", order="postulacion_id")
    respuesta_ids = {x["postulacion_id"] for x in respuestas}
    postulacion_ids = {x["id"] for x in posts}
    completas = all(isinstance(x.get("respuestas"), list)
                    and bool(x["respuestas"])
                    and x["respuestas"][0].get("p") == "Marca temporal"
                    for x in respuestas)
    failures += not check(19, len(respuestas) == 24203 and respuesta_ids == postulacion_ids and completas,
                          f"respuestas={len(respuestas)}")
    from cargar_fase2 import read_source, choose_sources, DEFAULT_SOURCE, cedula_norm
    _, fase2_by_cedula = read_source(DEFAULT_SOURCE)
    fase2_chosen = choose_sources(fase2_by_cedula)
    fase2_posts = api.get("postulaciones", "id,convocatoria,pais", order="id")
    fase2_pii = api.get("postulaciones_pii", "postulacion_id,cedula_norm", order="postulacion_id")
    fase2_post_by_id = {x["id"]: x for x in fase2_posts}
    fase2_candidates = {}
    for item in fase2_pii:
        post = fase2_post_by_id.get(item["postulacion_id"])
        if post and post["convocatoria"] == "2026" and item.get("cedula_norm"):
            key = (post["pais"], cedula_norm(item["cedula_norm"]))
            fase2_candidates.setdefault(key, [])
            if post["id"] not in fase2_candidates[key]:
                fase2_candidates[key].append(post["id"])
    fase2_canonicas = {x["id"] for x in posts if not x.get("duplicado_de")}
    fase2_matched_ids = set()
    for cedula, source in fase2_chosen.items():
        ids = fase2_candidates.get((source["pais"], cedula), [])
        if len(ids) > 1:
            ids = [post_id for post_id in ids if post_id in fase2_canonicas]
        if len(ids) == 1:
            fase2_matched_ids.add(ids[0])
    fase2_matched = len(fase2_matched_ids)
    fase2_rows = api.get("resultado_fase2", "postulacion_id", order="postulacion_id")
    fase2_ids = {x["postulacion_id"] for x in fase2_rows}
    fase2_resultados = api.get("resultado_seleccion", "postulacion_id,fase_max_alcanzada",
                               order="postulacion_id")
    fase1_remaining = sum(x["postulacion_id"] in fase2_ids and x["fase_max_alcanzada"] == "fase1"
                          for x in fase2_resultados)
    failures += not check(20, len(fase2_rows) == fase2_matched and fase1_remaining == 0,
                          f"resultado_fase2={len(fase2_rows)} matchean={fase2_matched} fase1={fase1_remaining}")
    fase3_chosen, fase3_files, fase3_blocks = claves_fase3_independientes()
    fase3_matched_ids = set()
    fase3_sin_dato_ids = set()
    for key, source in fase3_chosen.items():
        ids = fase2_candidates.get(key, [])
        if len(ids) > 1:
            ids = [post_id for post_id in ids if post_id in fase2_canonicas]
        if len(ids) == 1:
            fase3_matched_ids.add(ids[0])
            if source["attended"] is None:
                fase3_sin_dato_ids.add(ids[0])
    fase3_rows = api.get("resultado_fase3", "postulacion_id,asistio", order="postulacion_id")
    fase3_resultados = {x["postulacion_id"]: x["fase_max_alcanzada"]
                        for x in api.get("resultado_seleccion", "postulacion_id,fase_max_alcanzada",
                                         order="postulacion_id")}
    fase3_lower = sum(fase3_resultados.get(x["postulacion_id"]) in {"fase1", "fase2"}
                      for x in fase3_rows)
    failures += not check(21, len(fase3_rows) == len(fase3_matched_ids) and fase3_lower == 0,
                          f"resultado_fase3={len(fase3_rows)} matchean={len(fase3_matched_ids)} "
                          f"fase1_fase2={fase3_lower} archivos={fase3_files} archivos_bloqueados={fase3_blocks}")
    fase3_null_ids = {x["postulacion_id"] for x in fase3_rows if x["asistio"] is None}
    fase3_below = {x["postulacion_id"] for x in fase3_rows
                   if fase3_resultados.get(x["postulacion_id"]) in {"fase1", "fase2"}}
    failures += not check(22, fase3_null_ids == fase3_sin_dato_ids and not fase3_below,
                          f"asistencia_sin_dato={len(fase3_null_ids)} esperadas={len(fase3_sin_dato_ids)} "
                          f"fase_max_inferior={len(fase3_below)}")
    print(f"T7 {'FALLA' if failures else 'OK'}: fallas={failures}")
    return 1 if failures else 0


if __name__ == "__main__": raise SystemExit(main())

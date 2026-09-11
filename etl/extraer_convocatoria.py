"""Extrae las ocho fuentes de convocatoria a JSONL sin transformar sus columnas."""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

try:
    import truststore
    truststore.inject_into_ssl()
except ImportError:
    pass

EXPECTED = {
    ("2026", "CO"): 11356,
    ("2026", "EC"): 1430,
    ("2026", "PA"): 567,
    ("2026", "UY"): 219,
    ("2025", "CO"): 7977,
    ("2025", "CO-form1"): 90,
    ("2025", "EC"): 2213,
    ("2025", "UY"): 351,
}


def _texto(v: object) -> str:
    return "" if v is None else str(v)


def _escribir(fh, source: str, rows, convocatoria: str, pais: str, start_row: int = 2) -> int:
    n = 0
    for row_num, row in enumerate(rows, start=start_row):
        values = [_texto(v) for v in row]
        while values and values[-1] == "":
            values.pop()
        record = {str(k): v for k, v in enumerate(values)}
        record["_fuente"] = source
        record["_fila"] = row_num
        record["_convocatoria"] = convocatoria
        record["_pais"] = pais.split("-")[0]
        record["_pais_archivo"] = pais
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")
        n += 1
    return n


def _csv(path: Path, out: Path, convocatoria: str, pais: str) -> int:
    with path.open("r", encoding="utf-8-sig", newline="") as src, out.open("w", encoding="utf-8") as dst:
        reader = csv.reader(src)
        header = next(reader)
        keys: list[str] = []
        seen: dict[str, int] = {}
        for value in header:
            key = _texto(value)
            seen[key] = seen.get(key, 0) + 1
            keys.append(key if seen[key] == 1 else f"{key}__{seen[key]}")
        n = 0
        for row_num, row in enumerate(reader, start=2):
            record = {keys[i]: _texto(row[i]) if i < len(row) else "" for i in range(len(keys))}
            record.update({"_fuente": path.name, "_fila": row_num,
                           "_convocatoria": convocatoria, "_pais": pais.split("-")[0],
                           "_pais_archivo": pais})
            dst.write(json.dumps(record, ensure_ascii=False) + "\n")
            n += 1
        return n


def _xlsx(path: Path, sheet: str, out: Path, convocatoria: str, pais: str) -> int:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        # Preserve source headings by storing them as keys, not numeric positions.
        # Duplicate headings receive a deterministic suffix.
        keys: list[str] = []
        seen: dict[str, int] = {}
        for value in header:
            key = _texto(value)
            seen[key] = seen.get(key, 0) + 1
            keys.append(key if seen[key] == 1 else f"{key}__{seen[key]}")
        with out.open("w", encoding="utf-8") as dst:
            n = 0
            for row_num, row in enumerate(rows, start=2):
                values = list(row)
                if not any(value not in (None, "") for value in values):
                    continue
                record = {keys[i]: _texto(values[i]) if i < len(values) else "" for i in range(len(keys))}
                record.update({"_fuente": f"{path.name}::{sheet}", "_fila": row_num,
                               "_convocatoria": convocatoria, "_pais": pais})
                dst.write(json.dumps(record, ensure_ascii=False) + "\n")
                n += 1
            return n
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dir", type=Path, default=Path(r"C:\Users\EstudiantesJC\Downloads"))
    parser.add_argument("--salida", type=Path, default=Path(__file__).parent / "salida")
    args = parser.parse_args()
    args.salida.mkdir(parents=True, exist_ok=True)
    d = args.dir
    xlsx = d / "Convocatoria Fase 1.xlsx"
    sources = [
        ("2026", "CO", xlsx, "Respuestas Colombia", None),
        ("2026", "EC", xlsx, "Respuestas Ecuador", None),
        ("2026", "PA", xlsx, "Respuestas Panamá", None),
        ("2026", "UY", xlsx, "Respuestas Uruguay", None),
        ("2025", "CO", d / "COL - Fase #1 Jóvenes creaTIvos 2025 (respuestas) - INTOCABLE.csv", None, "raw_2025_CO"),
        ("2025", "CO-form1", d / "COL - Fase #1 Jóvenes creaTIvos 2025 (respuestas) - Respuestas de formulario 1.csv", None, "raw_2025_CO-form1"),
        ("2025", "EC", d / "ECU - Fase #1 Jóvenes creaTIvos 2025 (respuestas) - INTOCABLE.csv", None, "raw_2025_EC"),
        ("2025", "UY", d / "Fase #1 Jóvenes creaTIvos 2025 UY (respuestas) - INTOCABLE.csv", None, "raw_2025_UY"),
    ]
    failures = 0
    for convocatoria, pais, path, sheet, forced_name in sources:
        output_name = forced_name or f"raw_{convocatoria}_{pais}.jsonl"
        out = args.salida / f"{output_name}.jsonl" if forced_name else args.salida / output_name
        if not path.exists():
            print(f"FALLA {path}: no existe", file=sys.stderr)
            failures += 1
            continue
        try:
            n = _xlsx(path, sheet, out, convocatoria, pais) if sheet else _csv(path, out, convocatoria, pais)
            expected = EXPECTED[(convocatoria, pais)]
            status = "OK" if n == expected else "FALLA"
            print(f"{status} {out.name}: {n} filas (esperadas {expected})")
            if n != expected:
                failures += 1
        except Exception as exc:
            print(f"FALLA {path}: {exc}", file=sys.stderr)
            failures += 1
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())

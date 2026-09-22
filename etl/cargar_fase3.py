"""Carga los resultados no-PII de la ronda de entrevistas de Fase 3."""
from __future__ import annotations

import argparse
import csv
import os
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
import truststore
truststore.inject_into_ssl()
import requests

from cargar_convocatoria import BATCH, Supabase, cargar_env_local
from cruzar_canon import cedula_norm

ROOT = Path(__file__).parent
SOURCE_ROOT = Path(r"C:\Users\EstudiantesJC\Downloads\2026-Fase3")
EXTRACTED = SOURCE_ROOT / "extraido"
ZIP_SOURCE = SOURCE_ROOT / "drive-download-20260922T161234Z-1-001.zip"
DUPLICATES = ROOT / "salida" / "fase3_duplicados_fuente.csv"

COUNTRY_BY_SUFFIX = {
    "BAQ": "CO", "BOG": "CO", "CLO": "CO", "CTG": "CO", "MED": "CO",
    "ECU": "EC", "URY": "UY", "PAN": "PA",
}
RANK = {"fase1": 0, "fase2": 1, "fase3": 2, "entrevista": 3,
        "seleccionado": 4, "matriculado": 5}


def header_key(value) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return " ".join(text.casefold().split())


def email_key(value) -> str:
    return "" if value is None else str(value).strip().casefold()


def number(value):
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def find_header(ws, required: set[str]) -> tuple[int, dict[str, int]]:
    for row_number, values in enumerate(ws.iter_rows(values_only=True), start=1):
        keys = {header_key(value) for value in values}
        if required.issubset(keys):
            columns = {}
            for index, value in enumerate(values):
                key = header_key(value)
                if key and key not in columns:
                    columns[key] = index
            return row_number, columns
    raise RuntimeError(f"no se encontró encabezado con {sorted(required)} en {ws.title}")


def cell(row, columns: dict[str, int], name: str):
    index = columns.get(header_key(name))
    return None if index is None or index >= len(row) else row[index]


def parse_file(path: Path, country: str, panel: int) -> tuple[list[dict], list[str]]:
    blocks: list[str] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        general = workbook["GENERAL"]
        support_name = next((name for name in workbook.sheetnames if name.startswith("SOPORTE")), None)
        if support_name is None:
            return [], [f"{path.name}: falta hoja SOPORTE"]
        support = workbook[support_name]
        general_header_row, general_columns = find_header(general, {"id", "email"})
        support_header_row, support_columns = find_header(support, {"email", "presente"})
        required_general = {"id", "email", "grupo", "capitan", "jurado 1", "jurado 2",
                            "total", "puntaje", "disponibilidad"}
        if not required_general.issubset(general_columns):
            return [], [f"{path.name}: faltan columnas GENERAL {sorted(required_general - set(general_columns))}"]

        availability_index = general_columns["disponibilidad"]
        score_index = general_columns["puntaje"]
        header_values = next(general.iter_rows(min_row=general_header_row,
                                                max_row=general_header_row,
                                                values_only=True))
        loose_indices = [index for index in range(score_index + 1, availability_index)
                         if index < len(header_values) and header_values[index] is not None
                         and str(header_values[index]).strip().isdigit()]
        fatal_blocks: list[str] = []
        if len(loose_indices) != 1:
            fatal_blocks.append(f"{path.name}: columna numérica suelta no identificada de forma única")
        cumple_index = loose_indices[0] if len(loose_indices) == 1 else None

        support_by_email: dict[str, list[str]] = defaultdict(list)
        for row in support.iter_rows(min_row=support_header_row + 1, values_only=True):
            email = email_key(cell(row, support_columns, "email"))
            if not email:
                continue
            support_by_email[email].append(header_key(cell(row, support_columns, "presente")))

        source_rows: list[dict] = []
        zero_emails: Counter[str] = Counter()
        no_emails: Counter[str] = Counter()
        seen_cumple = set()
        indeterminate = 0
        for row_number, row in enumerate(general.iter_rows(min_row=general_header_row + 1, values_only=True),
                                         start=general_header_row + 1):
            raw_id = cell(row, general_columns, "id")
            email = email_key(cell(row, general_columns, "email"))
            if raw_id is None and not email:
                continue
            if raw_id is None or not email:
                indeterminate += 1
                continue
            try:
                cedula = str(int(raw_id))
            except (TypeError, ValueError, OverflowError):
                indeterminate += 1
                continue
            total = number(cell(row, general_columns, "total"))
            if total == 0:
                zero_emails[email] += 1
            if cumple_index is not None:
                seen_cumple.add(header_key(row[cumple_index])
                                if cumple_index < len(row) and row[cumple_index] is not None else "")
            support_rows = support_by_email.get(email, [])
            if len(support_rows) != 1:
                indeterminate += 1
                continue
            present = support_rows[0]
            if present == "no":
                no_emails[email] += 1
            if present not in {"si", "no"}:
                indeterminate += 1
                attended = None
            else:
                attended = present == "si"
            source_rows.append({
                "cedula": cedula, "pais": country, "email": email,
                "grupo": cell(row, general_columns, "grupo"), "panel": panel,
                "asistio": attended,
                "puntaje_capitan": number(cell(row, general_columns, "capitan")) if attended is True else None,
                "puntaje_jurado1": number(cell(row, general_columns, "jurado 1")) if attended is True else None,
                "puntaje_jurado2": number(cell(row, general_columns, "jurado 2")) if attended is True else None,
                "total": total,
                "puntaje_pct": number(cell(row, general_columns, "puntaje")) if attended is True else None,
                "archivo": path.name, "fila": row_number,
                "cumple": (str(row[cumple_index]).strip() if cumple_index is not None and cumple_index < len(row) else ""),
            })
        if zero_emails != no_emails:
            fatal_blocks.append(f"{path.name}: Total=0 y Presente=NO no coinciden por email")
        if seen_cumple != {"cumple"}:
            fatal_blocks.append(f"{path.name}: columna numérica suelta no es constante CUMPLE")
        if indeterminate:
            blocks.append(f"{path.name}: {indeterminate} filas con ID, email o asistencia indeterminada")
        blocks.extend(fatal_blocks)
        return source_rows, blocks
    finally:
        workbook.close()


def ensure_source() -> Path:
    if EXTRACTED.exists():
        return EXTRACTED
    if not ZIP_SOURCE.exists():
        raise RuntimeError(f"no existe fuente extraída ni ZIP: {EXTRACTED} / {ZIP_SOURCE}")
    EXTRACTED.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(ZIP_SOURCE) as archive:
        archive.extractall(EXTRACTED)
    return EXTRACTED


def read_source() -> tuple[list[dict], dict[tuple[str, str], list[dict]], list[str], int]:
    root = ensure_source()
    paths = sorted(root.glob("*/Tabla de Puntuación * PANEL *.xlsx"))
    if len(paths) != 48:
        raise RuntimeError(f"se esperaban 48 archivos Excel, encontrados {len(paths)}")
    rows: list[dict] = []
    by_key: dict[tuple[str, str], list[dict]] = defaultdict(list)
    audit_by_key: dict[tuple[str, str], list[dict]] = defaultdict(list)
    blocks: list[str] = []
    for path in paths:
        suffix = path.parent.name[-3:]
        country = COUNTRY_BY_SUFFIX.get(suffix)
        if country is None:
            blocks.append(f"{path.parent.name}: sufijo de país desconocido")
            continue
        match = re.search(r"PANEL\s+(\d+)", path.name, re.IGNORECASE)
        if not match:
            blocks.append(f"{path.name}: no se pudo leer número de panel")
            continue
        file_rows, file_blocks = parse_file(path, country, int(match.group(1)))
        rows.extend(file_rows)
        blocks.extend(file_blocks)
        file_fatal = any("Total=0 y Presente=NO" in block or
                         "no es constante CUMPLE" in block or
                         "columna numérica suelta" in block or
                         "faltan columnas" in block or
                         "falta hoja SOPORTE" in block
                         for block in file_blocks)
        for row in file_rows:
            row["archivo_bloqueado"] = file_fatal
            audit_by_key[(row["pais"], row["cedula"])].append(row)
        if not file_fatal:
            for row in file_rows:
                if row["asistio"] is not None:
                    by_key[(row["pais"], row["cedula"])].append(row)
    write_duplicate_report(audit_by_key)
    return rows, by_key, blocks, len(paths)


def write_duplicate_report(by_key: dict[tuple[str, str], list[dict]]) -> None:
    DUPLICATES.parent.mkdir(parents=True, exist_ok=True)
    with DUPLICATES.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["pais", "cedula_norm", "filas_fuente", "fuentes", "fuente_elegida",
                         "fila_elegida", "asistio_elegida", "puntaje_pct_elegido", "criterio"])
        for key, candidates in sorted(by_key.items()):
            if len(candidates) <= 1:
                continue
            blocked = [x for x in candidates if x.get("archivo_bloqueado") or x["asistio"] is None]
            if blocked:
                winner = None
                criterion = "bloqueado: archivo o asistencia no validada"
            else:
                winner = max(candidates, key=lambda item: (item["asistio"],
                                                            item["puntaje_pct"] is not None,
                                                            item["puntaje_pct"] or -1,
                                                            item["archivo"], item["fila"]))
                criterion = "asistio > puntaje_pct" if len({x["asistio"] for x in candidates}) > 1 else "puntaje_pct más alto"
            writer.writerow([key[0], key[1], len(candidates),
                             " | ".join(f'{x["archivo"]}:fila {x["fila"]}' for x in candidates),
                             winner["archivo"] if winner else "",
                             winner["fila"] if winner else "",
                             winner["asistio"] if winner else "",
                             winner["puntaje_pct"] if winner else "", criterion])


def choose_sources(by_key: dict[tuple[str, str], list[dict]]) -> dict[tuple[str, str], dict]:
    chosen: dict[tuple[str, str], dict] = {}
    for key, candidates in sorted(by_key.items()):
        if len(candidates) > 1:
            known = [x for x in candidates if x["asistio"] is not None]
            if len(known) != len(candidates):
                continue
            winner = max(candidates, key=lambda item: (item["asistio"],
                                                        item["puntaje_pct"] is not None,
                                                        item["puntaje_pct"] or -1,
                                                        item["archivo"], item["fila"]))
        else:
            winner = candidates[0]
        chosen[key] = winner
    return chosen


def fetch_all(api: Supabase, table: str, select: str, order: str = "id") -> list[dict]:
    result: list[dict] = []
    offset = 0
    while True:
        page = api.request("GET", table, params={"select": select, "order": order,
                                                  "limit": BATCH, "offset": offset}).json()
        result.extend(page)
        if len(page) < BATCH:
            return result
        offset += BATCH


def match_sources(api: Supabase, chosen: dict[tuple[str, str], dict]) -> tuple[list[dict], int, list[str]]:
    posts = fetch_all(api, "postulaciones", "id,convocatoria,pais")
    pii = fetch_all(api, "postulaciones_pii", "postulacion_id,cedula_norm", "postulacion_id")
    selections = fetch_all(api, "resultado_seleccion", "postulacion_id,duplicado_de", "postulacion_id")
    post_by_id = {row["id"]: row for row in posts}
    duplicado_de = {row["postulacion_id"]: row.get("duplicado_de") for row in selections}
    candidates: dict[tuple[str, str], list[str]] = defaultdict(list)
    for item in pii:
        post = post_by_id.get(item["postulacion_id"])
        if post and post["convocatoria"] == "2026" and item.get("cedula_norm"):
            key = (post["pais"], cedula_norm(item["cedula_norm"]))
            if post["id"] not in candidates[key]:
                candidates[key].append(post["id"])
    matched: list[dict] = []
    ambiguous: list[str] = []
    for key, source in chosen.items():
        ids = candidates.get(key, [])
        if len(ids) > 1:
            canonicas = [post_id for post_id in ids if not duplicado_de.get(post_id)]
            ids = canonicas if len(canonicas) == 1 else []
            if not ids:
                ambiguous.append(f"país={key[0]} candidatos={len(candidates.get(key, []))}")
        if len(ids) != 1:
            continue
        matched.append({"postulacion_id": ids[0], **{field: source[field] for field in (
            "grupo", "panel", "asistio", "puntaje_capitan", "puntaje_jurado1",
            "puntaje_jurado2", "total", "puntaje_pct")}})
    return matched, len(chosen) - len(matched), ambiguous


def management_query(sql: str) -> None:
    url = os.environ["CONV_SUPABASE_URL"]
    ref = url.split("//", 1)[1].split(".", 1)[0]
    response = requests.post(
        f"https://api.supabase.com/v1/projects/{ref}/database/query",
        headers={"Authorization": f"Bearer {os.environ['CONV_SUPABASE_ACCESS_TOKEN']}",
                 "Content-Type": "application/json"},
        json={"query": sql}, timeout=180,
    )
    if response.status_code >= 300:
        raise RuntimeError(f"actualización de fase3 rechazada: HTTP {response.status_code} {response.text[:500]}")


def write_rows(api: Supabase, matched: list[dict]) -> None:
    for start in range(0, len(matched), BATCH):
        api.upsert("resultado_fase3", matched[start:start + BATCH], "postulacion_id")
    management_query("""
        update resultado_seleccion rs
        set fase_max_alcanzada = case
          when case rs.fase_max_alcanzada
            when 'fase1' then 0 when 'fase2' then 1 when 'fase3' then 2
            when 'entrevista' then 3 when 'seleccionado' then 4
            when 'matriculado' then 5 else 0 end < 2
          then 'fase3' else rs.fase_max_alcanzada
        end
        where exists (
          select 1 from resultado_fase3 rf3
          where rf3.postulacion_id = rs.postulacion_id
        )
    """)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--escribir", action="store_true")
    args = parser.parse_args()
    cargar_env_local()
    url = os.environ.get("CONV_SUPABASE_URL")
    key = os.environ.get("CONV_SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        raise SystemExit("faltan credenciales CONV_SUPABASE")
    rows, by_key, blocks, file_count = read_source()
    chosen = choose_sources(by_key)
    api = Supabase(url, key)
    matched, unmatched, ambiguous = match_sources(api, chosen)
    print(f"Archivos: {file_count}")
    print(f"Filas crudas: {len(rows)}")
    unique_source_ids = len({(row["pais"], row["cedula"]) for row in rows})
    print(f"IDs únicos en fuente: {unique_source_ids}")
    print(f"IDs elegibles tras bloqueos: {len(chosen)}")
    if DUPLICATES.exists():
        with DUPLICATES.open(encoding="utf-8", newline="") as handle:
            duplicate_rows = list(csv.DictReader(handle))
        resolved_duplicates = sum(bool(row["fuente_elegida"]) for row in duplicate_rows)
    else:
        duplicate_rows, resolved_duplicates = [], 0
    print(f"Duplicados de fuente: {len(duplicate_rows)}; resueltos: {resolved_duplicates}; "
          f"bloqueados: {len(duplicate_rows) - resolved_duplicates}")
    print(f"Matchean: {len(matched)}")
    print(f"No matchean: {unmatched}")
    if blocks:
        print("Bloqueos:")
        for block in blocks:
            print(f"- {block}")
    if ambiguous:
        print(f"Cruces ambiguos sin canónica: {len(ambiguous)}")
    if args.escribir:
        write_rows(api, matched)
        print(f"Escritas resultado_fase3: {len(matched)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
